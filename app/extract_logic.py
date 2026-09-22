import re
from pptx import Presentation

SUB_RE = re.compile(r'^[IVXLC]+\.\s?\d')
MAIN_RE = re.compile(r'^[IVXLC]+\s*[–\-\.]\s*\S')

def slide_texts(slide):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append(sh.text_frame.text.strip())
    return out

def classify(texts):
    if not texts:
        return 'EMPTY'
    first = texts[0]
    joined = ' '.join(texts)
    if re.match(r'^LIÇÃO\s*\d+', first, re.I):
        return 'TITLE'
    if first.strip().upper() == 'INTRODUÇÃO':
        return 'INTRO_BUMPER'
    if first.strip().upper().startswith('CONCLUSÃO'):
        return 'CONCLUSAO_BUMPER'
    if first.strip().upper().startswith('REVISANDO'):
        return 'REVISANDO_BUMPER'
    if SUB_RE.match(first.strip()):
        return 'SUB_BUMPER'
    if MAIN_RE.match(first.strip()):
        return 'MAIN_BUMPER'
    return 'CONTENT'

def extract_indices(path):
    p = Presentation(path)
    slides = list(p.slides)
    classes = [classify(slide_texts(s)) for s in slides]

    try:
        title_idx = next(i for i,c in enumerate(classes) if c == 'TITLE')
    except StopIteration:
        raise ValueError(f'{path}: não encontrei um slide de título ("LIÇÃO NN")')
    try:
        intro_idx = next(i for i,c in enumerate(classes) if c == 'INTRO_BUMPER')
    except StopIteration:
        raise ValueError(f'{path}: não encontrei o slide "INTRODUÇÃO"')

    # intro content = slides after intro_idx until first MAIN_BUMPER
    i = intro_idx + 1
    intro_content = []
    while i < len(classes) and classes[i] != 'MAIN_BUMPER':
        intro_content.append(i)
        i += 1
    if i >= len(classes):
        raise ValueError(f'{path}: não encontrei nenhum bumper de tópico (ex: "I – ...") depois da introdução')

    main_bumpers = []
    while len(main_bumpers) < 3:
        if i >= len(classes):
            raise ValueError(f'{path}: encontrei só {len(main_bumpers)} bumpers de tópico (esperava 3 — I, II, III)')
        if classes[i] == 'MAIN_BUMPER':
            main_bumpers.append(i)
        i += 1

    return {
        'title': title_idx,
        'intro_bumper': intro_idx,
        'intro_content': intro_content,
        'main_bumpers': main_bumpers,
        'classes': classes,
    }

if __name__ == '__main__':
    import sys
    for f in ['Lição_01_-_Abraão_-_Seu_chamado_e_sua_jornada_de_fé.pptx',
              'Lição_02_-_A_Fé_de_Abrão_nas_promessas_de_Deus.pptx',
              'Lição_03_-_A_impaciência_na_espera_do_cumprimento_da_promessa.pptx']:
        r = extract_indices(f)
        print(f)
        print('  title:', r['title']+1, 'intro_bumper:', r['intro_bumper']+1,
              'intro_content:', [i+1 for i in r['intro_content']],
              'main_bumpers:', [i+1 for i in r['main_bumpers']])
